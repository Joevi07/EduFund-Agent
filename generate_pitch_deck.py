import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    # 16:9 Widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]

    # Color Palette
    BG_DARK = RGBColor(11, 16, 30)
    BG_CARD = RGBColor(15, 23, 42)
    TEXT_WHITE = RGBColor(248, 250, 252)
    TEXT_MUTED = RGBColor(148, 163, 184)
    CYAN = RGBColor(32, 168, 216)
    EMERALD = RGBColor(46, 204, 113)
    INDIGO = RGBColor(99, 102, 241)
    AMBER = RGBColor(243, 156, 18)
    CORAL = RGBColor(231, 76, 60)

    def add_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_DARK
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category_text="EDUFUND PITCH DECK"):
        # Category Pill
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(0.4))
        tf_cat = cat_box.text_frame
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = CYAN

        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.7), Inches(0.8))
        tf_title = title_box.text_frame
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(26)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_WHITE

    # ==================== SLIDE 1: Title Slide ====================
    slide1 = prs.slides.add_slide(blank_layout)
    add_bg(slide1)

    # Accent Card Shape
    card1 = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(1.5), Inches(10.333), Inches(4.5))
    card1.fill.solid()
    card1.fill.fore_color.rgb = BG_CARD
    card1.line.color.rgb = CYAN
    card1.line.width = Pt(2)

    tf1 = card1.text_frame
    tf1.word_wrap = True
    
    p1 = tf1.paragraphs[0]
    p1.text = "EduFund AI Agent OS"
    p1.font.size = Pt(40)
    p1.font.bold = True
    p1.font.color.rgb = CYAN
    p1.alignment = PP_ALIGN.CENTER

    p2 = tf1.add_paragraph()
    p2.text = "AI-Powered Education Funding Discovery, Financial Planning & Application Autopilot"
    p2.font.size = Pt(20)
    p2.font.color.rgb = TEXT_WHITE
    p2.alignment = PP_ALIGN.CENTER
    p2.space_before = Pt(20)

    p3 = tf1.add_paragraph()
    p3.text = "Bridging the Education Funding Gap with Multi-Agent Intelligence"
    p3.font.size = Pt(14)
    p3.font.color.rgb = EMERALD
    p3.alignment = PP_ALIGN.CENTER
    p3.space_before = Pt(30)

    # ==================== SLIDE 2: Problem Statement ====================
    slide2 = prs.slides.add_slide(blank_layout)
    add_bg(slide2)
    add_header(slide2, "The Core Problem in Education Funding")

    problems = [
        ("Fragmented Opportunities", "Funding is scattered across hundreds of portals, university sites, and grant listings.", CORAL),
        ("Eligibility Complexity", "Students waste time applying for opportunities with hidden restrictions or missed requirements.", AMBER),
        ("Repetitive Friction", "Filling out identical marksheets, income proofs, and essays creates severe application fatigue.", CORAL),
        ("No Financial Strategy", "Scholarship sites act as static directories without calculating actual out-of-pocket funding gaps.", AMBER)
    ]

    for idx, (p_title, p_desc, color) in enumerate(problems):
        row = idx // 2
        col = idx % 2
        left = Inches(0.8 + col * 5.9)
        top = Inches(1.8 + row * 2.6)

        box = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(5.6), Inches(2.3))
        box.fill.solid()
        box.fill.fore_color.rgb = BG_CARD
        box.line.color.rgb = color
        box.line.width = Pt(1.5)

        tf = box.text_frame
        tf.word_wrap = True

        p_h = tf.paragraphs[0]
        p_h.text = p_title
        p_h.font.size = Pt(18)
        p_h.font.bold = True
        p_h.font.color.rgb = color

        p_b = tf.add_paragraph()
        p_b.text = p_desc
        p_b.font.size = Pt(13)
        p_b.font.color.rgb = TEXT_WHITE
        p_b.space_before = Pt(10)

    # ==================== SLIDE 3: The Multi-Agent Solution ====================
    slide3 = prs.slides.add_slide(blank_layout)
    add_bg(slide3)
    add_header(slide3, "The EduFund Solution — 6 Autonomous Agents")

    agents = [
        ("1. Profile Agent", "Parses GPA, major, family income & budget.", INDIGO),
        ("2. Discovery Agent", "Queries curated funding repositories.", CYAN),
        ("3. Eligibility Agent", "Performs line-by-line criteria reasoning.", EMERALD),
        ("4. Planner Agent ⭐", "Solves Funding Gap with EV portfolio stack.", AMBER),
        ("5. Autopilot Agent 🔥", "AI Essay Studio & document verification.", INDIGO),
        ("6. Deadline Agent 🗺️", "Monitors closing dates & document reuse.", CORAL)
    ]

    for idx, (a_title, a_desc, color) in enumerate(agents):
        row = idx // 3
        col = idx % 3
        left = Inches(0.8 + col * 3.9)
        top = Inches(1.8 + row * 2.6)

        box = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(3.7), Inches(2.3))
        box.fill.solid()
        box.fill.fore_color.rgb = BG_CARD
        box.line.color.rgb = color
        box.line.width = Pt(1.5)

        tf = box.text_frame
        tf.word_wrap = True

        p_h = tf.paragraphs[0]
        p_h.text = a_title
        p_h.font.size = Pt(16)
        p_h.font.bold = True
        p_h.font.color.rgb = color

        p_b = tf.add_paragraph()
        p_b.text = a_desc
        p_b.font.size = Pt(12)
        p_b.font.color.rgb = TEXT_WHITE
        p_b.space_before = Pt(10)

    # ==================== SLIDE 4: FinTech Engine & Gap Calculation ====================
    slide4 = prs.slides.add_slide(blank_layout)
    add_bg(slide4)
    add_header(slide4, "FinTech Financial Engine & Gap Calculation")

    # Equation Box
    eq_box = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.7), Inches(1.8))
    eq_box.fill.solid()
    eq_box.fill.fore_color.rgb = BG_CARD
    eq_box.line.color.rgb = CYAN
    eq_box.line.width = Pt(2)

    tf_eq = eq_box.text_frame
    tf_eq.word_wrap = True

    pe1 = tf_eq.paragraphs[0]
    pe1.text = "THE EDUCATION FUNDING EQUATION"
    pe1.font.size = Pt(12)
    pe1.font.bold = True
    pe1.font.color.rgb = CYAN

    pe2 = tf_eq.add_paragraph()
    pe2.text = "Target Annual Education Cost  −  Confirmed Aid  =  Remaining Funding Gap"
    pe2.font.size = Pt(22)
    pe2.font.bold = True
    pe2.font.color.rgb = TEXT_WHITE
    pe2.alignment = PP_ALIGN.CENTER
    pe2.space_before = Pt(15)

    # Confidence Meter Box
    conf_box = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(3.9), Inches(11.7), Inches(3.0))
    conf_box.fill.solid()
    conf_box.fill.fore_color.rgb = BG_CARD
    conf_box.line.color.rgb = EMERALD
    conf_box.line.width = Pt(1.5)

    tf_conf = conf_box.text_frame
    tf_conf.word_wrap = True

    pc1 = tf_conf.paragraphs[0]
    pc1.text = "📊 Funding Confidence Meter Engine"
    pc1.font.size = Pt(18)
    pc1.font.bold = True
    pc1.font.color.rgb = EMERALD

    pc2 = tf_conf.add_paragraph()
    pc2.text = "Statistical Probability Algorithm (0 to 100% Score):"
    pc2.font.size = Pt(14)
    pc2.font.bold = True
    pc2.font.color.rgb = TEXT_WHITE
    pc2.space_before = Pt(10)

    pc3 = tf_conf.add_paragraph()
    pc3.text = "• 45% Strategy Gap Coverage Ratio  |  30% Avg Academic/Financial Match\n• 15% Document Verification Readiness  |  10% Time-to-Deadline Buffer"
    pc3.font.size = Pt(13)
    pc3.font.color.rgb = TEXT_MUTED
    pc3.space_before = Pt(8)

    # ==================== SLIDE 5: Expected Value Strategy Solver ====================
    slide5 = prs.slides.add_slide(blank_layout)
    add_bg(slide5)
    add_header(slide5, "Expected Value (EV) Strategy Portfolio Solver")

    card_ev1 = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.1))
    card_ev1.fill.solid()
    card_ev1.fill.fore_color.rgb = BG_CARD
    card_ev1.line.color.rgb = AMBER
    card_ev1.line.width = Pt(1.5)

    tf_ev1 = card_ev1.text_frame
    tf_ev1.word_wrap = True

    p_ev1 = tf_ev1.paragraphs[0]
    p_ev1.text = "Expected Value (EV) Formulation"
    p_ev1.font.size = Pt(18)
    p_ev1.font.bold = True
    p_ev1.font.color.rgb = AMBER

    p_ev2 = tf_ev1.add_paragraph()
    p_ev2.text = "EV = Funding Amount × Eligibility Match % × Urgency Weight"
    p_ev2.font.size = Pt(14)
    p_ev2.font.bold = True
    p_ev2.font.color.rgb = TEXT_WHITE
    p_ev2.space_before = Pt(15)

    p_ev3 = tf_ev1.add_paragraph()
    p_ev3.text = "• Prioritizes high-payoff, high-match opportunities.\n• Minimizes wasted effort on unlikely applications.\n• Solves optimal combinations to achieve 100% gap coverage."
    p_ev3.font.size = Pt(13)
    p_ev3.font.color.rgb = TEXT_MUTED
    p_ev3.space_before = Pt(15)

    # Simulator Side
    card_ev2 = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(1.8), Inches(5.6), Inches(5.1))
    card_ev2.fill.solid()
    card_ev2.fill.fore_color.rgb = BG_CARD
    card_ev2.line.color.rgb = CYAN
    card_ev2.line.width = Pt(1.5)

    tf_ev2 = card_ev2.text_frame
    tf_ev2.word_wrap = True

    p_sim1 = tf_ev2.paragraphs[0]
    p_sim1.text = "Real-Time Strategy Scenario Simulator"
    p_sim1.font.size = Pt(18)
    p_sim1.font.bold = True
    p_sim1.font.color.rgb = CYAN

    p_sim2 = tf_ev2.add_paragraph()
    p_sim2.text = "Interactive 'What-If' Parameter Controls:"
    p_sim2.font.size = Pt(14)
    p_sim2.font.bold = True
    p_sim2.font.color.rgb = TEXT_WHITE
    p_sim2.space_before = Pt(15)

    p_sim3 = tf_ev2.add_paragraph()
    p_sim3.text = "• 'What if family contribution increases by ₹15,000?'\n• 'What if a Work-Study stipend is added?'\n• Real-time solver re-calculates remaining gap and EV stack."
    p_sim3.font.size = Pt(13)
    p_sim3.font.color.rgb = TEXT_MUTED
    p_sim3.space_before = Pt(15)

    # ==================== SLIDE 6: Application Autopilot 🔥 ====================
    slide6 = prs.slides.add_slide(blank_layout)
    add_bg(slide6)
    add_header(slide6, "Application Autopilot & AI Essay Studio 🔥")

    ap_features = [
        ("Tailored Response Generation", "Synthesizes custom SOPs & essay responses tailored to student achievements and opportunity prompts.", INDIGO),
        ("AI Essay Refiner Presets", "Instant tone modifiers: '⚡ High Impact', '🎓 Academic Rigor', '✂️ Shorten (<100w)'.", CYAN),
        ("Human Safety Review Control", "Safety Guarantee: AI never submits financial documents without explicit student authorization.", EMERALD)
    ]

    for idx, (title, desc, color) in enumerate(ap_features):
        left = Inches(0.8)
        top = Inches(1.8 + idx * 1.7)

        box = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(11.7), Inches(1.4))
        box.fill.solid()
        box.fill.fore_color.rgb = BG_CARD
        box.line.color.rgb = color
        box.line.width = Pt(1.5)

        tf = box.text_frame
        tf.word_wrap = True

        p_h = tf.paragraphs[0]
        p_h.text = title
        p_h.font.size = Pt(16)
        p_h.font.bold = True
        p_h.font.color.rgb = color

        p_b = tf.add_paragraph()
        p_b.text = desc
        p_b.font.size = Pt(13)
        p_b.font.color.rgb = TEXT_WHITE
        p_b.space_before = Pt(6)

    # ==================== SLIDE 7: Evidence Checker & Document Auditor ====================
    slide7 = prs.slides.add_slide(blank_layout)
    add_bg(slide7)
    add_header(slide7, "AI Evidence Checker & Document Auditor 🔍")

    card_evid = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.1))
    card_evid.fill.solid()
    card_evid.fill.fore_color.rgb = BG_CARD
    card_evid.line.color.rgb = EMERALD
    card_evid.line.width = Pt(1.5)

    tf_evid = card_evid.text_frame
    tf_evid.word_wrap = True

    p_evid1 = tf_evid.paragraphs[0]
    p_evid1.text = "🔍 AI Essay Evidence Checker"
    p_evid1.font.size = Pt(18)
    p_evid1.font.bold = True
    p_evid1.font.color.rgb = EMERALD

    p_evid2 = tf_evid.add_paragraph()
    p_evid2.text = "• Cross-references essay statements against verified student profile facts.\n• Highlights validated claims in green ('✓ VALIDATED').\n• Flags unverified claims in amber ('⚠️ UNVERIFIED PROFILE')."
    p_evid2.font.size = Pt(13)
    p_evid2.font.color.rgb = TEXT_WHITE
    p_evid2.space_before = Pt(15)

    card_doc = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(1.8), Inches(5.6), Inches(5.1))
    card_doc.fill.solid()
    card_doc.fill.fore_color.rgb = BG_CARD
    card_doc.line.color.rgb = CYAN
    card_doc.line.width = Pt(1.5)

    tf_doc = card_doc.text_frame
    tf_doc.word_wrap = True

    p_doc1 = tf_doc.paragraphs[0]
    p_doc1.text = "📄 Document Verification Auditor"
    p_doc1.font.size = Pt(18)
    p_doc1.font.bold = True
    p_doc1.font.color.rgb = CYAN

    p_doc2 = tf_doc.add_paragraph()
    p_doc2.text = "• Audits attached mock PDFs (marksheets, income certificates, LORs).\n• Verifies document readiness against opportunity rules.\n• Flags missing signatures or unverified income dates."
    p_doc2.font.size = Pt(13)
    p_doc2.font.color.rgb = TEXT_WHITE
    p_doc2.space_before = Pt(15)

    # ==================== SLIDE 8: Reuse Map & Collision Detector ====================
    slide8 = prs.slides.add_slide(blank_layout)
    add_bg(slide8)
    add_header(slide8, "Document Reuse Map & Deadline Collisions 🗺️")

    card_reuse = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.1))
    card_reuse.fill.solid()
    card_reuse.fill.fore_color.rgb = BG_CARD
    card_reuse.line.color.rgb = CYAN
    card_reuse.line.width = Pt(1.5)

    tf_reuse = card_reuse.text_frame
    tf_reuse.word_wrap = True

    p_reuse1 = tf_reuse.paragraphs[0]
    p_reuse1.text = "🗺️ Document Reuse Mapping"
    p_reuse1.font.size = Pt(18)
    p_reuse1.font.bold = True
    p_reuse1.font.color.rgb = CYAN

    p_reuse2 = tf_reuse.add_paragraph()
    p_reuse2.text = "• Maps 1 uploaded document across multiple target applications.\n• Example: '1 Income Certificate satisfies 3 scholarship requirements!'\n• Reduces repetitive upload effort by 60%+."
    p_reuse2.font.size = Pt(13)
    p_reuse2.font.color.rgb = TEXT_WHITE
    p_reuse2.space_before = Pt(15)

    card_col = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(1.8), Inches(5.6), Inches(5.1))
    card_col.fill.solid()
    card_col.fill.fore_color.rgb = BG_CARD
    card_col.line.color.rgb = CORAL
    card_col.line.width = Pt(1.5)

    tf_col = card_col.text_frame
    tf_col.word_wrap = True

    p_col1 = tf_col.paragraphs[0]
    p_col1.text = "⏰ Deadline Collision Detector"
    p_col1.font.size = Pt(18)
    p_col1.font.bold = True
    p_col1.font.color.rgb = CORAL

    p_col2 = tf_col.add_paragraph()
    p_col2.text = "• Audits closing dates to detect application collisions within tight 3–7 day windows.\n• Highlights collision risk levels ('HIGH_COLLISION').\n• Recommends an optimized submission sequence."
    p_col2.font.size = Pt(13)
    p_col2.font.color.rgb = TEXT_WHITE
    p_col2.space_before = Pt(15)

    # ==================== SLIDE 9: Tech Stack ====================
    slide9 = prs.slides.add_slide(blank_layout)
    add_bg(slide9)
    add_header(slide9, "Technical Architecture & Tech Stack")

    tech_box = slide9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.1))
    tech_box.fill.solid()
    tech_box.fill.fore_color.rgb = BG_CARD
    tech_box.line.color.rgb = CYAN
    tech_box.line.width = Pt(1.5)

    tf_tech = tech_box.text_frame
    tf_tech.word_wrap = True

    pt1 = tf_tech.paragraphs[0]
    pt1.text = "🐍 Backend Architecture (Python FastAPI)"
    pt1.font.size = Pt(18)
    pt1.font.bold = True
    pt1.font.color.rgb = CYAN

    pt2 = tf_tech.add_paragraph()
    pt2.text = "• FastAPI REST API engine with Pydantic schemas and Uvicorn runner.\n• SQLite database persistence (edufund.db) for student profiles & applications.\n• Modular agent pipeline (Profile, Discovery, Eligibility, Planner, Autopilot, Deadline)."
    pt2.font.size = Pt(13)
    pt2.font.color.rgb = TEXT_WHITE
    pt2.space_before = Pt(10)

    pt3 = tf_tech.add_paragraph()
    pt3.text = "⚛️ Frontend UI (React + Vite)"
    pt3.font.size = Pt(18)
    pt3.font.bold = True
    pt3.font.color.rgb = EMERALD
    pt3.space_before = Pt(25)

    pt4 = tf_tech.add_paragraph()
    pt4.text = "• React single-page app built with Vite and Lucide React Icons.\n• Bento Grid Command Center & Collapsible Glass Sidebar.\n• Dual Currency Engine (₹ INR / $ USD toggle).\n• Interactive How-It-Works Slide-over Side Panel."
    pt4.font.size = Pt(13)
    pt4.font.color.rgb = TEXT_WHITE
    pt4.space_before = Pt(10)

    # ==================== SLIDE 10: Call to Action & Conclusion ====================
    slide10 = prs.slides.add_slide(blank_layout)
    add_bg(slide10)

    card10 = slide10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(1.5), Inches(10.333), Inches(4.5))
    card10.fill.solid()
    card10.fill.fore_color.rgb = BG_CARD
    card10.line.color.rgb = EMERALD
    card10.line.width = Pt(2)

    tf10 = card10.text_frame
    tf10.word_wrap = True

    p10_1 = tf10.paragraphs[0]
    p10_1.text = "Empowering Education Funding Equity"
    p10_1.font.size = Pt(36)
    p10_1.font.bold = True
    p10_1.font.color.rgb = EMERALD
    p10_1.alignment = PP_ALIGN.CENTER

    p10_2 = tf10.add_paragraph()
    p10_2.text = "EduFund turns static scholarship search into an autonomous, AI-guided financial strategy."
    p10_2.font.size = Pt(18)
    p10_2.font.color.rgb = TEXT_WHITE
    p10_2.alignment = PP_ALIGN.CENTER
    p10_2.space_before = Pt(20)

    p10_3 = tf10.add_paragraph()
    p10_3.text = "Experience Live Platform: http://127.0.0.1:5173"
    p10_3.font.size = Pt(16)
    p10_3.font.bold = True
    p10_3.font.color.rgb = CYAN
    p10_3.alignment = PP_ALIGN.CENTER
    p10_3.space_before = Pt(30)

    # Save Presentation
    pptx_path = os.path.join(os.path.dirname(__file__), "EduFund_Pitch_Deck.pptx")
    prs.save(pptx_path)
    print(f"Presentation saved to: {pptx_path}")

if __name__ == "__main__":
    create_presentation()
